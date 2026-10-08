import csv
import re
from collections import Counter
from pathlib import Path

from docx import Document
from pptx import Presentation


PROJECT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT / "outputs"
AUDIT_DIR = OUTPUT / "audits"
TERM_PATTERNS = {
    "oracle": re.compile(r"\boracle\w*\b", re.IGNORECASE),
    "adaptive": re.compile(r"\badaptive\w*\b", re.IGNORECASE),
    "optimal_or_optimum": re.compile(r"\boptim(?:al\w*|um\w*)\b", re.IGNORECASE),
    "saving": re.compile(r"\bsaving\w*\b", re.IGNORECASE),
    "regret": re.compile(r"\bregret\w*\b", re.IGNORECASE),
    "improvement": re.compile(r"\bimprov\w*\b", re.IGNORECASE),
    "material": re.compile(r"\bmaterial\w*\b", re.IGNORECASE),
    "redesign": re.compile(r"\bredesign\w*\b", re.IGNORECASE),
    "benefit": re.compile(r"\bbenefit\w*\b", re.IGNORECASE),
}


def docx_text(path: Path):
    document = Document(path)
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip():
            yield f"paragraph:{index}", paragraph.text.strip()
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                text = " | ".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )
                if text:
                    yield (
                        f"table:{table_index}:row:{row_index}:cell:{column_index}",
                        text,
                    )


def pptx_text(path: Path):
    presentation = Presentation(path)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes):
            text = getattr(shape, "text", "").strip()
            if text:
                yield f"slide:{slide_index}:shape:{shape_index}", text


def text_file(path: Path):
    for line_number, line in enumerate(
        path.read_text(encoding="utf-8").splitlines(), start=1
    ):
        if line.strip():
            yield f"line:{line_number}", line.strip()


def matched_terms(text: str) -> list[str]:
    return [
        term for term, pattern in TERM_PATTERNS.items() if pattern.search(text)
    ]


def main() -> None:
    artifacts = [
        (OUTPUT / "manuscript" / "manuscript_anonymized.docx", docx_text),
        (OUTPUT / "submission" / "cover_letter.docx", docx_text),
        (OUTPUT / "submission" / "highlights.txt", text_file),
        (OUTPUT / "figures" / "Figures_editable.pptx", pptx_text),
        (OUTPUT / "tables" / "Tables_editable.docx", docx_text),
        (OUTPUT / "tables" / "Tables_editable.pptx", pptx_text),
        (
            OUTPUT / "submission" / "graphical_abstract_editable.pptx",
            pptx_text,
        ),
    ]
    rows = []
    term_counts: Counter[str] = Counter()
    artifact_counts: Counter[str] = Counter()
    for path, extractor in artifacts:
        if not path.exists():
            raise FileNotFoundError(path)
        relative_path = path.relative_to(PROJECT).as_posix()
        for location, text in extractor(path):
            terms = matched_terms(text)
            if not terms:
                continue
            rows.append(
                {
                    "artifact": relative_path,
                    "location": location,
                    "matched_terms": ";".join(terms),
                    "text": text,
                }
            )
            artifact_counts[relative_path] += 1
            term_counts.update(terms)

    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    occurrence_path = AUDIT_DIR / "oracle_language_occurrences.csv"
    with occurrence_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["artifact", "location", "matched_terms", "text"],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)

    summary_path = AUDIT_DIR / "oracle_language_summary.csv"
    with summary_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["summary_type", "name", "count"])
        for name, count in sorted(artifact_counts.items()):
            writer.writerow(["artifact", name, count])
        for name, count in sorted(term_counts.items()):
            writer.writerow(["term", name, count])

    print(f"Wrote {len(rows)} audited language occurrences")


if __name__ == "__main__":
    main()
