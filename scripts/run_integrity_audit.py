import csv
import hashlib
import json
import os
from pathlib import Path

from docx import Document
from PIL import Image


PROJECT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT / "outputs"
REPORTS = PROJECT / "reports"
REPRODUCIBILITY_CRITICAL_SOURCE_IDS = {
    "AN-NC03-ISO216-SAMPLE",
    "SEL-BE03-UCI-DATA",
    "SEL-TR01-NYC-DATA",
    "SEL-UI01-OPSD-DATA",
    "SEL-UI09-NASA-POWER",
    "SEL-WK07-NASA-DATA",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_audit() -> tuple[list[str], list[str]]:
    failures = []
    notes = []
    missing_evidence = []
    require_full_archive = os.environ.get(
        "LEGACY_STANDARD_AUDIT_REQUIRE_FULL_ARCHIVE"
    ) == "1"
    with (PROJECT / "SOURCE_REGISTRY.csv").open(
        newline="", encoding="utf-8"
    ) as handle:
        rows = list(csv.DictReader(handle))
    registered_ids = {row["source_id"] for row in rows}
    missing_required_records = REPRODUCIBILITY_CRITICAL_SOURCE_IDS - registered_ids
    if missing_required_records:
        failures.append(
            "Missing reproducibility-critical source records: "
            + ", ".join(sorted(missing_required_records))
        )
    checked = 0
    for row in rows:
        path = PROJECT / row["raw_path"]
        if not path.exists():
            if (
                row["source_id"] in REPRODUCIBILITY_CRITICAL_SOURCE_IDS
                or require_full_archive
            ):
                failures.append(
                    f"Missing raw snapshot: {row['source_id']} ({row['raw_path']})"
                )
            else:
                missing_evidence.append(row["source_id"])
            continue
        checked += 1
        if int(row["file_size_bytes"]) != path.stat().st_size:
            failures.append(f"Size mismatch: {row['source_id']}")
        if row["sha256"] != sha256(path):
            failures.append(f"SHA-256 mismatch: {row['source_id']}")
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["completeness"]] = counts.get(row["completeness"], 0) + 1
    notes.append(
        f"Checked {checked} of {len(rows)} registered source snapshots; completeness labels: "
        + ", ".join(f"{key}={counts[key]}" for key in sorted(counts))
        + "."
    )
    if missing_evidence:
        notes.append(
            f"{len(missing_evidence)} evidence-only snapshots are not distributed in this "
            "clone; their provenance and expected hashes remain registered. Run "
            "`LEGACY_STANDARD_AUDIT_REQUIRE_FULL_ARCHIVE=1 make audit` against the "
            "preserved local archive for strict validation."
        )
    return failures, notes


def manuscript_audit() -> tuple[list[str], list[str]]:
    failures = []
    notes = []
    path = OUTPUT / "manuscript" / "manuscript_anonymized.docx"
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    try:
        abstract_index = paragraphs.index("Abstract")
        abstract = paragraphs[abstract_index + 1]
        abstract_words = len(abstract.split())
        if abstract_words > 250:
            failures.append(f"Abstract has {abstract_words} words; maximum is 250.")
        notes.append(f"Abstract word count: {abstract_words}.")
    except ValueError:
        failures.append("Abstract heading not found.")
    keyword_lines = [text for text in paragraphs if text.startswith("Keywords:")]
    if len(keyword_lines) != 1:
        failures.append("Exactly one keyword line is required.")
    else:
        keyword_count = len(keyword_lines[0].split(":", 1)[1].split(";"))
        if not 1 <= keyword_count <= 7:
            failures.append(f"Keyword count is {keyword_count}; required range is 1–7.")
        notes.append(f"Keyword count: {keyword_count}.")
    full_text = "\n".join(paragraphs)
    for label in ["Figure 1", "Figure 2", "Figure 3", "Table 1", "Table 2", "Table 3", "Table 4"]:
        if label not in full_text:
            failures.append(f"Missing manuscript citation/caption: {label}.")
    if "Declaration of generative AI" not in full_text:
        failures.append("Generative-AI disclosure is missing.")
    if "[To be supplied" in full_text:
        failures.append("Anonymized manuscript contains unresolved author placeholders.")
    references_index = paragraphs.index("References") if "References" in paragraphs else -1
    if references_index < 0:
        failures.append("References heading is missing.")
    else:
        references = [text for text in paragraphs[references_index + 1 :] if text]
        doi_count = sum("https://doi.org/" in reference for reference in references)
        notes.append(f"References: {len(references)}; DOI-linked references: {doi_count}.")
    return failures, notes


def submission_audit() -> tuple[list[str], list[str], list[str]]:
    failures = []
    notes = []
    actions = []
    required = [
        OUTPUT / "manuscript" / "title_page.docx",
        OUTPUT / "manuscript" / "manuscript_anonymized.docx",
        OUTPUT / "manuscript" / "manuscript_anonymized.pdf",
        OUTPUT / "supplement" / "supplement.docx",
        OUTPUT / "supplement" / "supplement.pdf",
        OUTPUT / "figures" / "Figures_editable.pptx",
        OUTPUT / "tables" / "Tables_editable.docx",
        OUTPUT / "tables" / "Tables_editable.pptx",
        OUTPUT / "submission" / "cover_letter.docx",
        OUTPUT / "submission" / "highlights.txt",
        OUTPUT / "submission" / "graphical_abstract.png",
        OUTPUT / "submission" / "TFSC_submission_package.zip",
    ]
    for path in required:
        if not path.exists() or path.stat().st_size == 0:
            failures.append(f"Missing or empty submission file: {path.relative_to(PROJECT)}")
    highlight_path = OUTPUT / "submission" / "highlights.txt"
    if highlight_path.exists():
        highlights = [
            line.removeprefix("• ").strip()
            for line in highlight_path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        if not 3 <= len(highlights) <= 5:
            failures.append(f"Highlight count is {len(highlights)}; required range is 3–5.")
        over = [item for item in highlights if len(item) > 85]
        if over:
            failures.append("At least one highlight exceeds 85 characters.")
        notes.append(
            "Highlight lengths: " + ", ".join(str(len(item)) for item in highlights) + "."
        )
    graphical = OUTPUT / "submission" / "graphical_abstract.png"
    if graphical.exists():
        width, height = Image.open(graphical).size
        if (width, height) != (1328, 531):
            failures.append(
                f"Graphical abstract is {width}×{height}; expected 1328×531 pixels."
            )
        notes.append(f"Graphical abstract dimensions: {width}×{height} pixels.")
    title_text = "\n".join(
        paragraph.text
        for paragraph in Document(
            OUTPUT / "manuscript" / "title_page.docx"
        ).paragraphs
    )
    if "[To be supplied" in title_text:
        actions.append("Supply author names, affiliations, and corresponding-author details.")
    cover_text = "\n".join(
        paragraph.text
        for paragraph in Document(
            OUTPUT / "submission" / "cover_letter.docx"
        ).paragraphs
    )
    if "[Corresponding author" in cover_text or "[Submission date]" in cover_text:
        actions.append("Complete the cover-letter date and corresponding-author signature.")
    declarations = (
        OUTPUT / "submission" / "declarations.txt"
    ).read_text(encoding="utf-8")
    if "must confirm" in declarations:
        actions.append("Confirm competing interests and funding declarations.")
    return failures, notes, actions


def values_audit() -> tuple[list[str], list[str]]:
    failures = []
    notes = []
    with (PROJECT / "MANUSCRIPT_VALUES.csv").open(
        newline="", encoding="utf-8"
    ) as handle:
        rows = list(csv.DictReader(handle))
    ids = [row["value_id"] for row in rows]
    if len(ids) != len(set(ids)):
        failures.append("MANUSCRIPT_VALUES.csv has duplicate value IDs.")
    for row in rows:
        output = PROJECT / row["analysis_output"]
        if not output.exists():
            failures.append(
                f"Missing source output for manuscript value {row['value_id']}: "
                f"{row['analysis_output']}"
            )
    ledger_rows = list(
        csv.DictReader(
            (PROJECT / "CLAIM_EVIDENCE_LEDGER.csv").open(
                newline="", encoding="utf-8"
            )
        )
    )
    unverified = [
        row["claim_id"]
        for row in ledger_rows
        if row["verification_status"] != "verified"
    ]
    if unverified:
        failures.append("Unverified claim-ledger entries: " + ", ".join(unverified))
    notes.append(f"Checked {len(rows)} canonical values and {len(ledger_rows)} claims.")
    return failures, notes


def write_report(
    path: Path,
    title: str,
    failures: list[str],
    notes: list[str],
    actions: list[str] | None = None,
) -> None:
    lines = [f"# {title}", "", f"Status: {'PASS' if not failures else 'FAIL'}", ""]
    lines.append("## Findings")
    if failures:
        lines.extend(f"- FAIL: {failure}" for failure in failures)
    else:
        lines.append("- No audit failures.")
    lines.extend(f"- {note}" for note in notes)
    if actions:
        lines.extend(["", "## Submission-time actions"])
        lines.extend(f"- {action}" for action in actions)
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def main() -> None:
    REPORTS.mkdir(parents=True, exist_ok=True)
    source_failures, source_notes = source_audit()
    manuscript_failures, manuscript_notes = manuscript_audit()
    submission_failures, submission_notes, actions = submission_audit()
    value_failures, value_notes = values_audit()
    write_report(
        REPORTS / "REPRODUCIBILITY_AUDIT.md",
        "Reproducibility audit",
        source_failures + value_failures,
        source_notes + value_notes,
    )
    write_report(
        REPORTS / "TFSC_COMPLIANCE_AUDIT.md",
        "TFSC compliance audit",
        manuscript_failures + submission_failures,
        manuscript_notes + submission_notes,
        actions,
    )
    all_failures = (
        source_failures + value_failures + manuscript_failures + submission_failures
    )
    write_report(
        REPORTS / "INTEGRITY_AUDIT.md",
        "Integrity audit",
        all_failures,
        source_notes + value_notes + manuscript_notes + submission_notes,
        actions,
    )
    summary = {
        "audit_version": "1.0.0",
        "failures": all_failures,
        "submission_time_actions": actions,
    }
    (OUTPUT / "submission" / "audit_summary.json").write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    if all_failures:
        raise SystemExit(f"Audit failed with {len(all_failures)} finding(s)")
    print("Integrity, reproducibility, and TFSC compliance audits passed")


if __name__ == "__main__":
    main()
