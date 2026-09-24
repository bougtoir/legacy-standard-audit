import subprocess
from pathlib import Path

import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH

from make_tables import tables
from publication_utils import (
    OUTPUT,
    add_figure,
    add_page_number,
    add_table,
    configured_document,
    display,
    interval,
    load_results,
    load_values,
)


MANUSCRIPT_DIR = OUTPUT / "manuscript"
FIGURE_DIR = OUTPUT / "figures"
TITLE = (
    "Legacy standards as frozen solutions to obsolete optimization problems: "
    "a framework for technological change and adaptive re-optimization"
)


def add_paragraph(document, text: str, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead) :])
    else:
        paragraph.add_run(text)


def add_equation(document, equation: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(equation)
    run.italic = True


def prior_art_references() -> list[str]:
    matrix = pd.read_csv(
        Path(__file__).resolve().parents[1] / "TFSC_prior_art_matrix.csv"
    )
    cited_ids = {
        "PA001",
        "PA002",
        "PA003",
        "PA004",
        "PA005",
        "PA006",
        "PA007",
        "PA008",
        "PA009",
        "PA010",
        "PA011",
        "PA012",
        "PA013",
        "PA014",
        "PA015",
        "PA016",
        "PA017",
        "PA018",
        "PA019",
        "PA020",
        "PA021",
        "PA022",
        "PA023",
        "PA024",
        "PA025",
        "PA026",
        "PA027",
        "PA028",
        "PA029",
        "PA030",
        "PA031",
        "PA032",
    }
    references = []
    for row in matrix[matrix["reference_id"].isin(cited_ids)].itertuples():
        authors = str(row.authors).replace(";", ",")
        references.append(
            f"{authors} ({int(row.year)}). {row.title}. {row.container}. "
            f"https://doi.org/{row.doi}"
        )
    institutional = [
        "Food and Agriculture Organization of the United Nations (2025). Crop evapotranspiration: Guidelines for computing crop water requirements, revised edition. FAO Irrigation and Drainage Paper 56 Rev.1.",
        "International Organization for Standardization (2007). ISO 216:2007 Writing paper and certain classes of printed matter—Trimmed sizes—A and B series, and indication of machine direction. https://www.iso.org/standard/36631.html",
        "International Organization for Standardization (2020). ISO 668:2020 Series 1 freight containers—Classification, dimensions and ratings. https://www.iso.org/standard/76912.html",
        "National Aeronautics and Space Administration (n.d.). C-MAPSS turbofan engine degradation simulation data. https://data.nasa.gov/docs/legacy/CMAPSSData.zip",
        "National Aeronautics and Space Administration POWER Project (n.d.). Daily agricultural meteorology for Davis, California, 2015–2024. https://power.larc.nasa.gov/",
        "New York City Department of Transportation (n.d.). Traffic Volume Counts (Historical). https://data.cityofnewyork.us/d/btm5-ppia",
        "Open Power System Data (2020). Time series, 60-minute single-index data package, version 2020-10-06. https://data.open-power-system-data.org/time_series/2020-10-06/",
        "UCI Machine Learning Repository (2024). Occupancy Detection [Dataset]. https://doi.org/10.24432/C5X01N",
        "United States Federal Highway Administration (2008). Traffic Signal Timing Manual. FHWA-HOP-08-024.",
    ]
    references.extend(institutional)
    return sorted(references, key=str.casefold)


def add_title_and_abstract(document, values) -> None:
    title = document.add_heading(TITLE, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(
        document,
        "Abstract",
        bold_lead="Abstract",
    )
    abstract = (
        "Persistent standards are often treated either as efficient coordination "
        "devices or as path-dependent constraints. Both interpretations can be "
        "premature when technologies have changed the information, computation, "
        "communication, and control conditions under which an incumbent was formed. "
        "We develop a Legacy Standard Audit that separates a historical optimum, a "
        "current frictionless optimum, and a transition-adjusted optimum. The protocol "
        f"was applied to a pre-specified universe of {display(values, 'DESIGN_CANDIDATES')} "
        f"candidates in {display(values, 'DESIGN_DOMAINS')} domains and a frozen set of "
        f"{display(values, 'DESIGN_FROZEN_CASES')} heterogeneous cases. Public inputs "
        "were snapshotted with provenance and checksums before confirmatory estimation. "
        "Oracle technical benchmarks identified material static mismatch in simulated "
        "condition-informed maintenance and reference-crop weather-adaptive irrigation; "
        "occupancy-responsive lighting and bounded price-aware load shifting remained "
        "small or uncertain after evidentiary and transition constraints were considered. "
        "A traffic-control sensitivity was insufficient for redesign, container "
        "compatibility could not be quantified, and the repeated-halving paper ratio "
        "returned the pre-specified analytic null. Because objectives and denominators "
        "differ, effects are not pooled and no population causal claim is made. The "
        "results show that technological change can justify re-audit without implying "
        "obsolescence: redesign requires case-specific transition-cost, safety, network, "
        "and regulatory evidence."
    )
    add_paragraph(document, abstract)
    add_paragraph(
        document,
        "Keywords: standards; path dependence; technological change; adaptive control; "
        "switching costs; reproducibility; technology governance",
    )


def add_introduction(document, values) -> None:
    document.add_heading("1. Introduction", level=1)
    add_paragraph(
        document,
        "Standards coordinate expectations, reduce interface uncertainty, and permit "
        "complementary investment. Those same benefits can make an incumbent persistent "
        "after the problem environment changes. Research on dominant designs and "
        "technological trajectories explains how technologies stabilize and later face "
        "discontinuities (Utterback and Abernathy, 1975; Dosi, 1982; Abernathy and Clark, "
        "1985; Tushman and Anderson, 1986; Henderson and Clark, 1990). Research on "
        "increasing returns, compatibility, and infrastructure lock-in explains why "
        "persistence can remain privately or socially rational (Farrell and Saloner, "
        "1985; Arthur, 1989; Liebowitz and Margolis, 1995; Pierson, 2000; Unruh, 2000). "
        "The analytical error is "
        "therefore not persistence itself, but inferring either efficiency or "
        "inefficiency from persistence without reconstructing the decision problem.",
    )
    add_paragraph(
        document,
        "Digital technologies sharpen this problem. Cheap sensing, computation, "
        "communication, and adaptive control can relax information constraints that once "
        "favored fixed schedules, fixed intervals, or static allocations. Digital "
        "innovation also changes architectures and organizing logics rather than merely "
        "improving a component (Yoo et al., 2010; Nambisan et al., 2017). Yet "
        "real-time information does not erase compatibility, investment, safety, "
        "regulation, or collective-action costs. Standards scholarship accordingly "
        "shows that dominance depends on installed bases, participation, governance, and "
        "the interaction between standardization and innovation (Allen and Sriram, 2000; "
        "Suarez, 2004; Wiegmann et al., 2017; van de Kaa et al., 2018; Blind et al., "
        "2022, 2023).",
    )
    add_paragraph(
        document,
        "This article contributes a pre-specified measurement and governance protocol. "
        "It does not offer a new general theory of lock-in. Instead, it operationalizes "
        "a question that spans standards economics, technological transitions, and "
        "adaptive systems: when a technology changes the feasible information or control "
        "set, how should an incumbent standard be re-audited without selecting only "
        "striking examples or forcing a redesign conclusion? The protocol requires "
        "separate estimates of static and transition-adjusted regret, preserves null and "
        "insufficient-evidence cases, and forbids pooling outcomes with incomparable "
        "objectives.",
    )
    add_paragraph(
        document,
        f"We screened {display(values, 'DESIGN_CANDIDATES')} candidates across "
        f"{display(values, 'DESIGN_DOMAINS')} domains, froze "
        f"{display(values, 'DESIGN_FROZEN_CASES')} cases before confirmatory estimation, "
        f"and grounded the positioning in {display(values, 'DESIGN_VERIFIED_PRIOR_ART')} "
        "exact-DOI-verified prior works. The empirical cases are demonstrations of a "
        "present-day audit, not complete reconstructions of every historical optimum. "
        "That distinction is central to the interpretation.",
    )


def add_theory(document) -> None:
    document.add_heading("2. Theory and hypotheses", level=1)
    add_paragraph(
        document,
        "Path dependence can arise through increasing returns and coordination, but "
        "historical persistence alone does not establish market failure. Critical work "
        "on escape from lock-in and technological trajectories emphasizes that incumbent "
        "systems may be difficult to displace even when alternatives improve (Cowan and "
        "Hultén, 1996; Dolfsma and Leydesdorff, 2009; Heinrich, 2014). Sociological and "
        "organizational accounts add that standards are enacted through practices and "
        "institutions rather than imposed as purely technical rules (Orlikowski, 1992; "
        "Timmermans and Epstein, 2010). Exploration and exploitation create a related "
        "governance tension: reassessment has option value, but continual change can "
        "destroy coordination value (March, 1991).",
    )
    add_paragraph(
        document,
        "Transition research frames system change as a co-evolution of technology, "
        "institutions, infrastructure, and user practices (Geels, 2002, 2004, 2005). "
        "Governance experiments can reveal transition conditions, but their success is "
        "contextual (Bos and Brown, 2012). Standards participation and policy further "
        "shape whose objectives and constraints enter a redesign process (Choung et al., "
        "2012; Blind and von Laer, 2022). Consumer switching research likewise shows that "
        "comparative technical value is only one component of adoption (Kamolsook et al., "
        "2019).",
    )
    add_paragraph(
        document,
        "The framework yields five falsifiable expectations. First, fixed rules created "
        "under information scarcity should show larger static mismatch where reliable "
        "real-time control becomes feasible. Second, adaptive control should improve the "
        "specified objective only when information and control gains exceed error and "
        "implementation costs. Third, non-technological changes such as demand, prices, "
        "demography, and regulation may explain apparent mismatch and must be modeled or "
        "named. Fourth, network and switching costs can rationally preserve an incumbent "
        "despite positive frictionless regret. Fifth, some inherited standards should "
        "remain near-optimal for their specified objective and appear as null controls.",
    )


def add_framework(document, table_specs) -> None:
    document.add_heading("3. The Legacy Standard Audit", level=1)
    add_paragraph(
        document,
        "Let the incumbent be x_inc. The historical optimum x_h* minimizes loss under "
        "the historical feasible set and parameters; the current frictionless optimum "
        "x_c* minimizes current loss without switching costs; and the transition-adjusted "
        "optimum x_T* adds switching, compatibility, and coordination costs. Figure 1 "
        "shows the resulting decision sequence.",
    )
    add_equation(document, "x_h* = arg min(x ∈ X_h) L_h(x; θ_h)")
    add_equation(document, "x_c* = arg min(x ∈ X_c) L_c(x; θ_c)")
    add_equation(
        document,
        "x_T* = arg min(x ∈ X_c) [L_c(x; θ_c) + C_switch(x_inc, x)]",
    )
    add_equation(
        document,
        "R_static = L_c(x_inc; θ_c) − L_c(x_c*; θ_c)",
    )
    add_figure(
        document,
        FIGURE_DIR / "Figure_1_audit_framework.png",
        "Figure 1. Legacy Standard Audit. Technological change triggers reassessment, "
        "but redesign requires transition-adjusted evidence.",
    )
    add_paragraph(
        document,
        "Table 1 defines the five audit stages. A case can show positive static regret "
        "and still retain the incumbent if discounted benefits do not exceed conversion, "
        "coordination, stranded-asset, safety, or regulatory costs. Conversely, a null "
        "case is informative because it demonstrates that persistence need not indicate "
        "technological inertia.",
    )
    _, headers, rows, caption = table_specs[0]
    add_table(document, headers, rows, caption)


def add_methods(document, values, table_specs) -> None:
    document.add_heading("4. Data and pre-specified methods", level=1)
    add_paragraph(
        document,
        "Candidate selection preceded confirmatory estimation. The registry covered the "
        "built environment, transport, automotive interfaces, work, education, "
        "healthcare, packaging, paper and digital legacies, utilities, household and "
        "demographic systems, information-scarcity systems, and null controls. Screening "
        "used historical evidence, technology relevance, decision clarity, public-data "
        "availability, identifiability, counterfactual feasibility, transition-cost "
        "feasibility, source quality, and comparability. Expected direction, magnitude, "
        "and statistical significance were excluded from selection. The frozen roles "
        "were four primary cases, one supplementary case, one network control, and one "
        "analytic null control.",
    )
    add_paragraph(
        document,
        "Every public input was saved as a local raw snapshot and registered with its "
        "URL, identifier, version, access time, retrieval conditions, file size, SHA-256 "
        "hash, terms, and completeness. Raw inputs were not overwritten. Analysis "
        "configuration, random seed, and bootstrap settings were fixed in a machine-"
        "readable configuration. Temporal uncertainty used circular block resampling; "
        "the C-MAPSS engine was the independent simulation unit.",
    )
    add_paragraph(
        document,
        "Classifications were operational decision labels rather than significance tests. "
        "BE03 required the lower uncertainty bound to exceed "
        f"{display(values, 'THRESHOLD_BE03')} with no labeled occupied false-off events; "
        "WK07 required a positive lower regret bound and relative regret above "
        f"{display(values, 'THRESHOLD_WK07')}; UI09 required the lower normalized-regret "
        f"bound to exceed {display(values, 'THRESHOLD_UI09')}. UI01 remained uncertain "
        "because response and equilibrium were not identified, while TR01 and NC01 "
        "required incumbent and transition evidence unavailable in the registered sources. "
        "The thresholds organize the frozen audit and are not universal welfare criteria.",
    )
    add_paragraph(
        document,
        "BE03 compares a weekday 07:00–19:00 nominal lighting schedule with an oracle "
        f"occupancy policy using a {display(values, 'BE03_HOLD')}-minute hold. WK07 uses "
        "NASA C-MAPSS FD001 simulated run-to-failure trajectories to compare a "
        "training-selected fixed replacement age with perfect remaining-life condition "
        "information. UI01 conserves daily energy while shifting a bounded fraction of "
        "the observed German-Luxembourg load profile toward lower observed day-ahead "
        "prices; prices are not re-equilibrated. UI09 calculates FAO-56 reference "
        "evapotranspiration from NASA POWER weather and compares a training-period "
        "calendar application with perfect-weather weekly requirements. TR01 is limited "
        "to a directional demand-allocation proxy. NC01 and NC03 test, respectively, an "
        "unquantified compatibility case and an exact geometric null. Source attribution "
        "is UCI Machine Learning Repository (2024) for BE03, National Aeronautics and "
        "Space Administration (n.d.) for WK07, Open Power System Data (2020) for UI01, "
        "Food and Agriculture Organization of the United Nations (2025) and National "
        "Aeronautics and Space Administration POWER Project (n.d.) for UI09, United "
        "States Federal Highway Administration (2008) and New York City Department of "
        "Transportation (n.d.) for TR01, and International Organization for "
        "Standardization (2007, 2020) for the controls.",
    )
    add_paragraph(
        document,
        "Table 2 lists the analysis unit, comparator, and primary threat to interpretation "
        "for each frozen case. These limitations are part of the result rather than "
        "post-estimation caveats.",
    )
    _, headers, rows, caption = table_specs[1]
    add_table(document, headers, rows, caption)


def add_results(document, values, table_specs) -> None:
    document.add_heading("5. Results", level=1)
    add_paragraph(
        document,
        f"BE03 retained {display(values, 'BE03_OBSERVATIONS')} valid minute-observations "
        "after removing two embedded repeated header rows. The oracle occupancy policy "
        f"reduced nominal on-time by {display(values, 'BE03_ENERGY_REDUCTION')} "
        f"(95% block interval {interval(values, 'BE03_ENERGY_REDUCTION')}) while both "
        "policies had zero occupied false-off observations under the retained labels. "
        "The result was classified as small or uncertain because schedule choice strongly "
        "changed the estimate and deployment costs and sensor errors were unavailable.",
    )
    add_paragraph(
        document,
        f"WK07 selected a fixed replacement age of {display(values, 'WK07_FIXED_AGE')} "
        "cycles on the training simulation. In the evaluation simulation, the fixed "
        f"policy cost rate was {display(values, 'WK07_FIXED_RATE')} and the oracle "
        f"condition rate was {display(values, 'WK07_CONDITION_RATE')}, yielding static "
        f"regret {display(values, 'WK07_STATIC_REGRET')} normalized cost per cycle "
        f"(95% interval {interval(values, 'WK07_STATIC_REGRET')}) and relative regret "
        f"{display(values, 'WK07_RELATIVE_REGRET')}. This is a material oracle benchmark, "
        "not observed fleet savings.",
    )
    add_paragraph(
        document,
        f"UI01 retained {display(values, 'UI01_COMPLETE_DAYS')} complete days. With "
        f"{display(values, 'UI01_FLEXIBLE_FRACTION')} of each hourly load available for "
        "bounded intraday reallocation, the technical procurement saving was "
        f"EUR {display(values, 'UI01_SAVINGS_PER_MWH')} per MWh of baseline load "
        f"(95% block interval {interval(values, 'UI01_SAVINGS_PER_MWH')}), equal to "
        f"{display(values, 'UI01_COST_REDUCTION')} of baseline procurement cost. The "
        "small estimate, fixed-price assumption, and absence of customer response support "
        "a small-or-uncertain classification.",
    )
    add_paragraph(
        document,
        f"UI09 used {display(values, 'UI09_TRAINING_WEEKS')} training and "
        f"{display(values, 'UI09_EVALUATION_WEEKS')} evaluation weeks. The calendar "
        f"application was {display(values, 'UI09_CALENDAR_APPLICATION')} mm/week, compared "
        f"with an adaptive mean of {display(values, 'UI09_ADAPTIVE_APPLICATION')} mm/week. "
        f"Normalized objective regret was {display(values, 'UI09_OBJECTIVE_REGRET')} "
        f"(95% block interval {interval(values, 'UI09_OBJECTIVE_REGRET')}). The result "
        "remained positive across pre-specified loss weights, but it is a reference-crop "
        "and perfect-weather benchmark with omitted agronomic constraints.",
    )
    add_paragraph(
        document,
        f"TR01 retained {display(values, 'TR01_EVALUATION_DAYS')} evaluation days and "
        f"produced a {display(values, 'TR01_PROXY_REDUCTION')} reduction in a directional "
        f"delay proxy (95% block interval {interval(values, 'TR01_PROXY_REDUCTION')}). "
        "Because the data do not establish an intersection, incumbent timing plan, "
        "pedestrian constraints, queues, saturation flows, or safety outcomes, the case "
        "remained insufficient evidence. NC01 also remained unquantified. NC03 returned "
        f"the exact positive repeated-halving ratio {display(values, 'NC03_OPTIMAL_RATIO')} "
        "with numerically zero loss for that narrow objective.",
    )
    add_paragraph(
        document,
        "Figure 2 reports primary-case benchmarks on separate scales to prevent false "
        "comparability. Table 3 reports the corresponding classifications.",
    )
    add_figure(
        document,
        FIGURE_DIR / "Figure_2_primary_case_benchmarks.png",
        "Figure 2. Frozen primary-case technical benchmarks. Each panel uses a "
        "case-specific objective and scale; estimates are not pooled.",
    )
    _, headers, rows, caption = table_specs[2]
    add_table(document, headers, rows, caption)


def add_synthesis(document, values, table_specs) -> None:
    document.add_heading("6. Cross-case synthesis and falsification", level=1)
    add_paragraph(
        document,
        f"Two of {display(values, 'CROSS_PRIMARY_CASES')} primary cases met the "
        "pre-specified material-oracle label, while the remaining primary cases were "
        "small or uncertain. That count is descriptive. Removing either material case "
        "reduced the count to one; excluding simulated C-MAPSS evidence also left one. "
        "The cross-case result therefore does not establish a stable population rate of "
        "mismatch.",
    )
    add_paragraph(
        document,
        "Figure 3 demonstrates why frozen null and insufficient-evidence cases matter. "
        "The outcomes span all intended categories rather than converging on a redesign "
        "narrative. The strongest evidence is conditional: inexpensive information and "
        "control can create measurable static regret under a specified objective, but "
        "transition-adjusted recommendations remain unidentified without implementation, "
        "network, safety, and coordination costs. Real-time foresight can improve "
        "preparedness in dynamic networks (Weber et al., 2015), but it cannot substitute "
        "for case-specific institutional evidence.",
    )
    add_figure(
        document,
        FIGURE_DIR / "Figure_3_case_classifications.png",
        "Figure 3. Classifications of all pre-specified frozen cases. Categories are "
        "decision labels, not hypothesis-test outcomes.",
    )
    _, headers, rows, caption = table_specs[3]
    add_table(document, headers, rows, caption)


def add_discussion(document) -> None:
    document.add_heading("7. Discussion", level=1)
    add_paragraph(
        document,
        "The main contribution is procedural. A standard should be audited as the output "
        "of a decision problem, not as an old object. The procedure begins by specifying "
        "the historical objective and constraints, identifies which constraints a "
        "technology plausibly changed, estimates a current technical counterfactual, and "
        "then asks whether transition costs reverse the ranking. This ordering prevents "
        "two symmetric errors: celebrating persistence as evidence of efficiency and "
        "treating positive static regret as sufficient evidence for redesign.",
    )
    add_paragraph(
        document,
        "The cases illustrate three technology mechanisms. Occupancy detection and "
        "weather data reduce the cost of sensing current state. Prognostics and interval "
        "metering reduce the cost of computing and communicating condition-dependent "
        "actions. Responsive control changes the feasible frequency of adjustment. Yet "
        "the same cases show why technological capability is not welfare evidence. "
        "Sensor errors, response costs, market equilibrium, crop biology, safety, and "
        "installed-base compatibility can dominate a frictionless gain.",
    )
    add_paragraph(
        document,
        "For technology governance, the implication is to institutionalize periodic "
        "audits rather than automatic sunset rules. The audit trigger should be a "
        "measurable change in information, control, production, or coordination "
        "conditions. The output should be a transparent classification with an explicit "
        "evidence gap. Regulators and standards bodies can then commission the missing "
        "transition evidence before altering a coordinated system. Multi-mode "
        "standardization and participation research suggests that this process should "
        "include affected implementers and complementors rather than treat the technical "
        "counterfactual as dispositive (Wiegmann et al., 2017; Blind and von Laer, 2022).",
    )
    add_paragraph(
        document,
        "The empirical results also discipline the novelty claim. Existing research "
        "already explains lock-in, standards battles, transition governance, digital "
        "innovation, and switching behavior. The added value here is a reproducible "
        "protocol that freezes case selection, separates static from transition-adjusted "
        "regret, retains negative controls, and records when available data are "
        "insufficient. This methodological synthesis is narrower than a general theory "
        "but more falsifiable than a collection of historical examples.",
    )


def add_limitations_conclusion(document) -> None:
    document.add_heading("8. Limitations", level=1)
    add_paragraph(
        document,
        "First, the analyses do not fully reconstruct the historical optimum for every "
        "case. They are frozen demonstrations of the current-audit stages. Second, four "
        "comparators are oracle or proxy bounds; none identifies a realized causal "
        "deployment effect. Third, transition-adjusted regret is not numerically "
        "estimated because compatible implementation-cost data are absent. Fourth, the "
        "case set is purposive and heterogeneous, so neither the frequency nor magnitude "
        "of mismatch generalizes to legacy standards as a population. Fifth, uncertainty "
        "intervals quantify sampling variation under the retained model but do not "
        "capture all structural uncertainty. Sixth, the largest normalized result, UI09, "
        "depends on a normative shortfall penalty and omits crop and soil processes.",
    )
    add_paragraph(
        document,
        "Future work should prospectively apply the audit to a second frozen set, acquire "
        "case-specific switching and implementation costs, and validate technical "
        "comparators against deployed controls. A particularly valuable extension would "
        "estimate one full transition-adjusted case from incumbent reconstruction through "
        "realized implementation.",
    )
    document.add_heading("9. Conclusion", level=1)
    add_paragraph(
        document,
        "Technological change can make a fixed standard reassessable without making it "
        "obsolete. The Legacy Standard Audit converts that distinction into a "
        "pre-specified workflow: reconstruct the decision problem, identify the changed "
        "technology constraint, estimate static regret, incorporate transition costs, and "
        "retain null or insufficient-evidence outcomes. In the frozen demonstrations, "
        "some oracle technical benchmarks were material, others were small, one analytic "
        "standard remained exactly aligned with its narrow objective, and a high-network "
        "compatibility case could not be quantified. The defensible conclusion is not "
        "that inherited standards should be replaced. It is that standards governance "
        "should make re-audit reproducible, conditional, and explicit about the evidence "
        "still required for change.",
    )


def add_end_matter(document) -> None:
    document.add_heading("Data and code availability", level=1)
    add_paragraph(
        document,
        "All analysis code, frozen selection records, source registry, checksums, and "
        "generated values are provided in the accompanying repository. Public raw data "
        "are retrieved from registered sources and retained as immutable local snapshots. "
        "Copyrighted ISO sample material is not redistributed; its provenance and local "
        "checksum are recorded. The one-command build regenerates analyses, figures, "
        "tables, manuscript files, supplement, and audits.",
    )
    document.add_heading("Declaration of generative AI and AI-assisted technologies", level=1)
    add_paragraph(
        document,
        "During preparation, the authors used Devin, an AI software-engineering assistant "
        "built by Cognition AI, to support code generation, document assembly, and "
        "consistency checks. Human authors are responsible for source verification, "
        "analysis choices, numerical results, interpretation, originality, and the final "
        "submitted text. The AI system is not an author.",
    )
    document.add_heading("Declarations", level=1)
    add_paragraph(document, "Funding: No project-specific funding statement supplied.")
    add_paragraph(document, "Competing interests: No statement supplied; authors must confirm before submission.")
    add_paragraph(document, "Ethics: The study uses public secondary data and simulations; no human participants were recruited.")
    document.add_heading("References", level=1)
    for reference in prior_art_references():
        add_paragraph(document, reference)


def build_title_page() -> None:
    document = configured_document()
    title = document.add_heading(TITLE, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    add_paragraph(document, "Article type: Full-length research article")
    add_paragraph(document, "Authors: [To be supplied by the submitting author]")
    add_paragraph(document, "Affiliations: [To be supplied]")
    add_paragraph(document, "Corresponding author: [Name, postal address, and email to be supplied]")
    add_paragraph(document, "Word count: generated manuscript; verify in final editorial check")
    add_paragraph(document, "Tables: 4 main; 1 supplementary")
    add_paragraph(document, "Figures: 3 main; 1 supplementary")
    add_paragraph(
        document,
        "Author details are intentionally not inferred from repository or account metadata.",
    )
    document.save(MANUSCRIPT_DIR / "title_page.docx")


def convert_to_pdf(path: Path) -> None:
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(path.parent),
            str(path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )


def main() -> None:
    MANUSCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    values = load_values()
    results = load_results()
    table_specs = tables(values, results)
    document = configured_document()
    add_page_number(document)
    add_title_and_abstract(document, values)
    add_introduction(document, values)
    add_theory(document)
    add_framework(document, table_specs)
    add_methods(document, values, table_specs)
    add_results(document, values, table_specs)
    add_synthesis(document, values, table_specs)
    add_discussion(document)
    add_limitations_conclusion(document)
    add_end_matter(document)
    manuscript_path = MANUSCRIPT_DIR / "manuscript_anonymized.docx"
    document.save(manuscript_path)
    build_title_page()
    convert_to_pdf(manuscript_path)
    print("Wrote anonymized manuscript, title page, and manuscript PDF")


if __name__ == "__main__":
    main()
