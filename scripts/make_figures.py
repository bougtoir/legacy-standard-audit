import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyBboxPatch
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from publication_utils import OUTPUT, display, interval, load_results, load_values


FIGURE_DIR = OUTPUT / "figures"
BLUE = "#176B87"
TEAL = "#2A9D8F"
GOLD = "#E9C46A"
RED = "#C8553D"
GRAY = "#6B7280"


def save_figure(fig: plt.Figure, name: str) -> None:
    fig.savefig(
        FIGURE_DIR / name,
        dpi=300,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(fig)


def figure_audit_framework() -> None:
    fig, ax = plt.subplots(figsize=(11, 5.7))
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 6)
    ax.axis("off")
    boxes = [
        (0.3, 3.7, 2.0, 1.1, "Historical problem", "Objective, information,\nconstraints, coordination"),
        (3.0, 3.7, 2.0, 1.1, "Technology change", "Sensing, computation,\ncommunication, control"),
        (5.7, 3.7, 2.0, 1.1, "Current static audit", "Incumbent versus\nfrictionless comparator"),
        (8.4, 3.7, 2.2, 1.1, "Transition audit", "Switching, safety,\nnetwork and policy costs"),
    ]
    for x, y, width, height, title, subtitle in boxes:
        patch = FancyBboxPatch(
            (x, y),
            width,
            height,
            boxstyle="round,pad=0.08",
            linewidth=1.4,
            edgecolor=BLUE,
            facecolor="#EAF4F8",
        )
        ax.add_patch(patch)
        ax.text(x + width / 2, y + 0.72, title, ha="center", va="center", weight="bold")
        ax.text(x + width / 2, y + 0.30, subtitle, ha="center", va="center", fontsize=9)
    for start, end in [(2.3, 3.0), (5.0, 5.7), (7.7, 8.4)]:
        ax.annotate(
            "",
            xy=(end, 4.25),
            xytext=(start, 4.25),
            arrowprops={"arrowstyle": "->", "lw": 1.8, "color": GRAY},
        )
    outcomes = [
        (0.8, 1.25, 2.7, "Material static mismatch", RED),
        (4.15, 1.25, 2.7, "Small or uncertain mismatch", GOLD),
        (7.5, 1.25, 2.7, "No mismatch / insufficient evidence", TEAL),
    ]
    for x, y, width, text, color in outcomes:
        patch = FancyBboxPatch(
            (x, y),
            width,
            0.8,
            boxstyle="round,pad=0.06",
            linewidth=1.2,
            edgecolor=color,
            facecolor="white",
        )
        ax.add_patch(patch)
        ax.text(x + width / 2, y + 0.4, text, ha="center", va="center", weight="bold")
    ax.text(
        5.5,
        0.45,
        "Redesign is supported only if transition-adjusted benefits remain positive.",
        ha="center",
        va="center",
        fontsize=11,
        weight="bold",
    )
    save_figure(fig, "Figure_1_audit_framework.png")


def figure_primary_benchmarks(values: dict[str, dict[str, str]]) -> None:
    panels = [
        ("BE03", "Energy reduction", "BE03_ENERGY_REDUCTION", "Fraction"),
        ("WK07", "Static regret", "WK07_STATIC_REGRET", "Normalized cost/cycle"),
        ("UI01", "Procurement savings", "UI01_SAVINGS_PER_MWH", "EUR/MWh baseline"),
        ("UI09", "Objective regret", "UI09_OBJECTIVE_REGRET", "Fraction"),
    ]
    fig, axes = plt.subplots(2, 2, figsize=(10, 7))
    for ax, (case_id, title, value_id, unit) in zip(axes.ravel(), panels):
        row = values[value_id]
        estimate = float(row["estimate"])
        lower = float(row["lower"])
        upper = float(row["upper"])
        margin = max((upper - lower) * 1.3, abs(estimate) * 0.18)
        ax.errorbar(
            estimate,
            0,
            xerr=np.array([[estimate - lower], [upper - estimate]]),
            fmt="o",
            color=BLUE,
            ecolor=BLUE,
            capsize=5,
            markersize=8,
        )
        ax.axvline(0, color="#A0A0A0", lw=1)
        ax.set_xlim(min(0, lower - margin), upper + margin)
        ax.set_yticks([])
        ax.set_xlabel(unit)
        ax.set_title(f"{case_id}: {title}", loc="left", weight="bold")
        ax.text(
            0.02,
            0.82,
            f"{display(values, value_id)} (95% block/bootstrap interval "
            f"{interval(values, value_id)})",
            transform=ax.transAxes,
            fontsize=9,
        )
        ax.spines[["top", "right", "left"]].set_visible(False)
    fig.suptitle(
        "Frozen primary-case technical benchmarks (case-specific scales)",
        weight="bold",
        fontsize=13,
    )
    fig.text(
        0.5,
        0.01,
        "Panels are not a pooled effect size; comparators are oracle or partial-equilibrium bounds.",
        ha="center",
        fontsize=9,
    )
    fig.tight_layout(rect=(0, 0.04, 1, 0.95))
    save_figure(fig, "Figure_2_primary_case_benchmarks.png")


def figure_case_classifications(results: dict[str, dict[str, object]]) -> None:
    cases = ["BE03", "WK07", "UI01", "UI09", "TR01", "NC01", "NC03"]
    columns = ["Material oracle", "Small/uncertain", "Insufficient", "Narrow null"]
    matrix = np.zeros((len(cases), len(columns)))
    for index, case_id in enumerate(cases):
        classification = str(results[case_id]["classification"])
        if classification.startswith("material_mismatch"):
            matrix[index, 0] = 1
        elif classification == "small_or_uncertain_mismatch":
            matrix[index, 1] = 1
        elif classification == "insufficient_evidence":
            matrix[index, 2] = 1
        else:
            matrix[index, 3] = 1
    fig, ax = plt.subplots(figsize=(9.5, 5.2))
    ax.imshow(matrix, cmap="Blues", vmin=0, vmax=1, aspect="auto")
    ax.set_xticks(range(len(columns)), columns, rotation=18, ha="right")
    ax.set_yticks(range(len(cases)), cases)
    for row in range(matrix.shape[0]):
        for column in range(matrix.shape[1]):
            ax.text(
                column,
                row,
                "●" if matrix[row, column] else "",
                ha="center",
                va="center",
                color="white",
                fontsize=16,
            )
    ax.set_title("Pre-specified cases retained across heterogeneous outcomes", weight="bold")
    ax.set_xlabel("Classification is a decision label, not a hypothesis test")
    fig.tight_layout()
    save_figure(fig, "Figure_3_case_classifications.png")


def sensitivity_rows(values: dict[str, dict[str, str]], case_id: str):
    rows = [
        row
        for value_id, row in values.items()
        if value_id.startswith(f"SENS_{case_id}_")
    ]
    return rows


def figure_sensitivity(values: dict[str, dict[str, str]]) -> None:
    cases = ["BE03", "WK07", "UI01", "UI09"]
    fig, axes = plt.subplots(2, 2, figsize=(11, 8))
    for ax, case_id in zip(axes.ravel(), cases):
        rows = sensitivity_rows(values, case_id)
        labels = [row["notes"].split(": ", 1)[1] for row in rows]
        estimates = [float(row["estimate"]) for row in rows]
        positions = np.arange(len(rows))
        ax.plot(positions, estimates, marker="o", color=BLUE)
        ax.set_xticks(positions, labels, rotation=45, ha="right", fontsize=8)
        ax.set_title(case_id, loc="left", weight="bold")
        ax.set_ylabel(rows[0]["unit"] if rows else "")
        ax.grid(axis="y", alpha=0.25)
    fig.suptitle("Pre-specified model sensitivity by case", weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    save_figure(fig, "Figure_S1_model_sensitivity.png")


def add_text_box(slide, x, y, width, height, text, fill, font_size=16, bold=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(23, 107, 135)
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return shape


def build_editable_pptx(
    values: dict[str, dict[str, str]],
    results: dict[str, dict[str, object]],
) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title_layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(title_layout)
    add_text_box(slide, 0.5, 0.25, 12.3, 0.65, "Figure 1. Legacy Standard Audit", (255, 255, 255), 24, True)
    labels = [
        "Historical problem\nObjective • constraints • information",
        "Technology change\nSensing • computation • control",
        "Static audit\nIncumbent vs frictionless comparator",
        "Transition audit\nSwitching • network • safety costs",
    ]
    for index, label in enumerate(labels):
        add_text_box(slide, 0.4 + index * 3.2, 2.0, 2.7, 1.3, label, (234, 244, 248), 14, True)
        if index < 3:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(3.0 + index * 3.2),
                Inches(2.45),
                Inches(0.6),
                Inches(0.35),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(107, 114, 128)
            arrow.line.fill.background()
    add_text_box(
        slide,
        2.7,
        5.25,
        8.0,
        0.8,
        "Recommend change only when transition-adjusted benefit remains positive",
        (255, 255, 255),
        18,
        True,
    )

    slide = presentation.slides.add_slide(title_layout)
    add_text_box(slide, 0.5, 0.25, 12.3, 0.65, "Figure 2. Primary-case benchmarks", (255, 255, 255), 24, True)
    benchmark_ids = [
        ("BE03", "BE03_ENERGY_REDUCTION"),
        ("WK07", "WK07_STATIC_REGRET"),
        ("UI01", "UI01_SAVINGS_PER_MWH"),
        ("UI09", "UI09_OBJECTIVE_REGRET"),
    ]
    for index, (case_id, value_id) in enumerate(benchmark_ids):
        x = 0.7 + (index % 2) * 6.2
        y = 1.3 + (index // 2) * 2.8
        add_text_box(slide, x, y, 5.5, 0.55, case_id, (234, 244, 248), 18, True)
        add_text_box(
            slide,
            x,
            y + 0.75,
            5.5,
            1.15,
            f"Estimate: {display(values, value_id)}\n95% interval: {interval(values, value_id)}\n{values[value_id]['unit']}",
            (255, 255, 255),
            15,
        )

    slide = presentation.slides.add_slide(title_layout)
    add_text_box(slide, 0.5, 0.25, 12.3, 0.65, "Figure 3. Frozen-case classifications", (255, 255, 255), 24, True)
    for row, case_id in enumerate(["BE03", "WK07", "UI01", "UI09", "TR01", "NC01", "NC03"]):
        add_text_box(slide, 0.8, 1.15 + row * 0.78, 1.2, 0.55, case_id, (234, 244, 248), 14, True)
        add_text_box(
            slide,
            2.2,
            1.15 + row * 0.78,
            9.8,
            0.55,
            str(results[case_id]["classification"]).replace("_", " "),
            (255, 255, 255),
            14,
        )

    slide = presentation.slides.add_slide(title_layout)
    add_text_box(slide, 0.5, 0.25, 12.3, 0.65, "Figure S1. Model sensitivity", (255, 255, 255), 24, True)
    for index, case_id in enumerate(["BE03", "WK07", "UI01", "UI09"]):
        rows = sensitivity_rows(values, case_id)
        text = "\n".join(
            f"{row['notes'].split(': ', 1)[1]}: {float(row['estimate']):.4g}"
            for row in rows
        )
        x = 0.7 + (index % 2) * 6.2
        y = 1.15 + (index // 2) * 3.0
        add_text_box(slide, x, y, 5.5, 0.55, case_id, (234, 244, 248), 17, True)
        add_text_box(slide, x, y + 0.65, 5.5, 2.0, text, (255, 255, 255), 10)

    presentation.save(FIGURE_DIR / "Figures_editable.pptx")


def main() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.titlecolor": "#1F2937",
            "axes.labelcolor": "#374151",
            "text.color": "#1F2937",
        }
    )
    values = load_values()
    results = load_results()
    figure_audit_framework()
    figure_primary_benchmarks(values)
    figure_case_classifications(results)
    figure_sensitivity(values)
    build_editable_pptx(values, results)
    print("Wrote 4 figures and editable PPTX")


if __name__ == "__main__":
    main()
