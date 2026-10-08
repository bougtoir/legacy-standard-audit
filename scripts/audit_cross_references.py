"""Audit figure/table cross-references and journal-facing supplement residue."""

import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

from publication_utils import (
    FIGURE_ORDER,
    OUTPUT,
    TABLE_ORDER,
    fig,
    figure_filename,
    tab,
)

PROJECT = OUTPUT.parent
REPORT = PROJECT / "reports" / "CROSS_REFERENCE_AUDIT.md"
MANUSCRIPT = OUTPUT / "manuscript" / "manuscript_anonymized.docx"
ARCHIVE = OUTPUT / "submission" / "TFSC_submission_package_FINAL.zip"
LABEL = re.compile(r"\b(Figure|Table) (\d+)\b")
SUPPLEMENT_POINTERS = re.compile(
    r"\b(?:Figure|Table|Section)\s+S\d+\b|\bS[1-9]\.\s|\bAppendix\b|\bsupplement\b|"
    r"\bSupplement\b|\bsupplementary (?:material|file|information|table|figure|"
    r"section|method|data|result)s?\b",
    re.IGNORECASE,
)
ROLE_LABEL = re.compile(r"\bsupplementary case\b", re.IGNORECASE)
BLACK = {None, "000000"}


def collect(document: Document):
    body, captions = [], []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text
        if paragraph.style.name == "Caption":
            match = LABEL.match(text)
            captions.append((index, match.group(0) if match else text))
        else:
            for match in LABEL.finditer(text):
                body.append((index, match.group(0)))
    return body, captions


def main() -> None:
    document = Document(MANUSCRIPT)
    body, captions = collect(document)
    failures: list[str] = []
    expected = {"Figure": [fig(key) for key, _ in FIGURE_ORDER],
                "Table": [tab(key) for key in TABLE_ORDER]}
    caption_labels = [label for _, label in captions]
    cited = [label for _, label in body]
    first_index = {}
    for index, label in body:
        first_index.setdefault(label, index)
    for kind, labels in expected.items():
        if [label for label in caption_labels if label.startswith(kind)] != labels:
            failures.append(f"{kind} captions out of order or missing: {caption_labels}")
        first = [label for label in dict.fromkeys(cited) if label.startswith(kind)]
        if first != labels:
            failures.append(f"{kind} first-citation order {first} != {labels}")
        for label in labels:
            caption_index = next((i for i, c in captions if c == label), None)
            if label not in first_index:
                failures.append(f"{label} is never cited in body text")
            elif caption_index is not None and first_index[label] > caption_index:
                failures.append(f"{label} is first cited after its caption")
    dangling = sorted(set(cited) - set(caption_labels), key=cited.index)
    if dangling:
        failures.append(f"Dangling citations: {dangling}")
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    pointers = sorted({m.group(0) for text in texts for m in SUPPLEMENT_POINTERS.finditer(
        ROLE_LABEL.sub("", text))})
    if pointers:
        failures.append(f"Supplement pointers remain: {pointers}")
    role_mentions = sum(len(ROLE_LABEL.findall(text)) for text in texts)
    if any("**" in text or "`" in text or re.search(r"(?<![\w])\*(?!\s)", text) for text in texts):
        failures.append("Markdown emphasis or code markup found in manuscript text")
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                runs.extend(run for paragraph in cell.paragraphs for run in paragraph.runs)
    if any("\n" in run.text or "\r" in run.text for run in runs):
        failures.append("Runs contain hard line breaks")
    if any(run.font.name not in {None, "Times New Roman"} for run in runs):
        failures.append("Runs override the Times New Roman font")
    colors = {str(run.font.color.rgb) if run.font.color.rgb else None for run in runs}
    if not colors <= BLACK:
        failures.append(f"Non-black run colors: {colors}")
    styles = document.styles
    for name in ["Normal", "Heading 1", "Heading 2", "Title", "Caption"]:
        color = styles[name].font.color.rgb
        if color is not None and str(color) != "000000":
            failures.append(f"Style {name} is not black")
    if styles["Normal"].font.size.pt != 12 or styles["Caption"].font.size.pt != 10:
        failures.append("Normal must be 12 pt and Caption 10 pt")
    omml = document.element.body.findall(".//" + qn("m:oMath"))
    if len(omml) < 4:
        failures.append(f"Expected editable OMML equations; found {len(omml)}")
    with zipfile.ZipFile(ARCHIVE) as archive:
        names = archive.namelist()
    if any(re.search(r"supplement|_S\d", name, re.IGNORECASE) for name in names):
        failures.append("Submission archive contains supplementary files")
    for key, _ in FIGURE_ORDER:
        if figure_filename(key) not in names:
            failures.append(f"Archive lacks {figure_filename(key)}")
    if (OUTPUT / "supplement").exists():
        failures.append("Stale outputs/supplement directory exists")

    lines = [
        "# Cross-reference audit",
        "",
        "Generated by scripts/audit_cross_references.py from the built manuscript and ZIP.",
        "",
        "| Object | First body citation (paragraph) | Caption (paragraph) | Body citations |",
        "|---|---:|---:|---:|",
    ]
    for label in expected["Figure"] + expected["Table"]:
        caption_index = next((i for i, c in captions if c == label), "missing")
        lines.append(
            f"| {label} | {first_index.get(label, 'never')} | {caption_index} | "
            f"{cited.count(label)} |"
        )
    lines += [
        "",
        f"- Editable OMML math objects: {len(omml)}.",
        f"- Case-role label 'supplementary case' (TR01 frozen role, not a file pointer): "
        f"{role_mentions} occurrence(s).",
        f"- Archive entries: {len(names)}; supplementary files: none." if not any(
            "supplement" in name.lower() for name in names) else "- Archive has supplement.",
        "",
        f"Status: {'FAIL' if failures else 'PASS'}",
    ]
    lines += [f"- FAIL: {failure}" for failure in failures]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if failures:
        print("\n".join(failures))
        sys.exit(1)
    print("Cross-reference audit passed")


if __name__ == "__main__":
    main()
