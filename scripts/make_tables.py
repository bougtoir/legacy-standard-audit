import csv

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from publication_utils import (
    OUTPUT,
    add_table,
    configured_document,
    display,
    interval,
    load_results,
    load_values,
)


TABLE_DIR = OUTPUT / "tables"


def write_csv(name: str, headers: list[str], rows: list[list[object]]) -> None:
    with (TABLE_DIR / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)


def tables(values, results):
    table1_headers = ["Audit stage", "Required evidence", "Decision output"]
    table1_rows = [
        [
            "Historical reconstruction",
            "Objective, information set, feasible set, coordination conditions",
            "Historically defensible incumbent problem",
        ],
        [
            "Technology-mediated change",
            "Sensing, computation, communication, automation, adaptive control",
            "Changed parameter or feasible-set claim",
        ],
        [
            "Static audit",
            "Incumbent and contemporary comparator under a common objective",
            "Case-specific frictionless regret",
        ],
        [
            "Transition audit",
            "Switching, compatibility, safety, regulation, network effects",
            "Transition-adjusted regret or break-even condition",
        ],
        [
            "Governance",
            "Frozen cases, uncertainty, null controls, falsification",
            "Material, uncertain, null, or insufficient classification",
        ],
    ]
    table2_headers = [
        "Case",
        "Role",
        "Persisted input",
        "Analysis unit",
        "Contemporary comparator",
        "Critical limitation",
    ]
    table2_rows = [
        [
            "BE03",
            "Primary",
            "UCI office occupancy",
            f"{display(values, 'BE03_OBSERVATIONS')} minute-observations",
            f"Oracle occupancy control, {display(values, 'BE03_HOLD')}-minute hold",
            "No sensor error, power, commissioning, or installation cost",
        ],
        [
            "WK07",
            "Primary",
            "NASA C-MAPSS FD001 simulation",
            f"{display(values, 'WK07_EVALUATION_UNITS')} evaluation units",
            "Perfect remaining-life condition policy",
            "Simulated trajectories; no observed maintenance program",
        ],
        [
            "UI01",
            "Primary",
            "OPSD load and day-ahead prices",
            f"{display(values, 'UI01_COMPLETE_DAYS')} complete days",
            f"{display(values, 'UI01_FLEXIBLE_FRACTION')} bounded load shift",
            "Fixed prices; no tariff assignment or customer response",
        ],
        [
            "UI09",
            "Primary",
            "NASA POWER weather and FAO-56",
            f"{display(values, 'UI09_EVALUATION_WEEKS')} evaluation weeks",
            "Perfect-weather reference-crop requirement",
            "No crop coefficient, soil storage, efficiency, or yield model",
        ],
        [
            "TR01",
            "Supplementary",
            "NYC directional traffic counts",
            f"{display(values, 'TR01_EVALUATION_DAYS')} evaluation days",
            "Demand-proportional directional allocation proxy",
            "No verified intersection or deployed timing plan",
        ],
        [
            "NC01",
            "Network control",
            "ISO container-standard metadata",
            "Not estimable",
            "Break-even identity only",
            "No fleet, conversion, or network-cost data",
        ],
        [
            "NC03",
            "Analytic null",
            "ISO 216 geometric principle",
            "Exact analytic objective",
            r"Positive solution r = √2",
            "Applies only to repeated-halving similarity",
        ],
    ]
    table3_headers = [
        "Case",
        "Pre-specified metric",
        "Estimate",
        "95% interval",
        "Classification",
    ]
    table3_rows = [
        [
            "BE03",
            "Nominal lighting-energy reduction",
            display(values, "BE03_ENERGY_REDUCTION"),
            interval(values, "BE03_ENERGY_REDUCTION"),
            display(values, "BE03_CLASSIFICATION"),
        ],
        [
            "WK07",
            "Static regret, normalized cost/cycle",
            display(values, "WK07_STATIC_REGRET"),
            interval(values, "WK07_STATIC_REGRET"),
            display(values, "WK07_CLASSIFICATION"),
        ],
        [
            "UI01",
            "Procurement saving, EUR/MWh baseline",
            display(values, "UI01_SAVINGS_PER_MWH"),
            interval(values, "UI01_SAVINGS_PER_MWH"),
            display(values, "UI01_CLASSIFICATION"),
        ],
        [
            "UI09",
            "Normalized objective regret",
            display(values, "UI09_OBJECTIVE_REGRET"),
            interval(values, "UI09_OBJECTIVE_REGRET"),
            display(values, "UI09_CLASSIFICATION"),
        ],
        [
            "TR01",
            "Delay-proxy reduction",
            display(values, "TR01_PROXY_REDUCTION"),
            interval(values, "TR01_PROXY_REDUCTION"),
            display(values, "TR01_CLASSIFICATION"),
        ],
        [
            "NC01",
            "Static/transition regret",
            "Not estimated",
            "Not estimated",
            display(values, "NC01_CLASSIFICATION"),
        ],
        [
            "NC03",
            "Repeated-halving loss",
            display(values, "NC03_SELF_SIMILARITY_LOSS"),
            "Exact analytic result",
            display(values, "NC03_CLASSIFICATION"),
        ],
    ]
    falsification = pd.read_csv(
        OUTPUT / "analysis" / "cross_case_falsification.csv"
    )
    table4_headers = [
        "Test",
        "Omitted case",
        "Primary retained",
        "Material oracle",
        "Uncertain/insufficient",
        "Interpretation",
    ]
    table4_rows = [
        [
            row.test.replace("_", " "),
            row.omitted_case if pd.notna(row.omitted_case) else "",
            int(row.primary_cases_retained),
            int(row.material_oracle_mismatches),
            int(row.insufficient_or_uncertain),
            row.interpretation,
        ]
        for row in falsification.itertuples()
    ]
    sensitivity_rows = []
    for value_id, row in values.items():
        if value_id.startswith("SENS_"):
            sensitivity_rows.append(
                [
                    row["case_id"],
                    row["notes"],
                    f"{float(row['estimate']):.6g}",
                    row["unit"],
                ]
            )
    return [
        ("Table_1_audit_framework.csv", table1_headers, table1_rows, "Table 1. Legacy Standard Audit stages."),
        ("Table_2_case_design.csv", table2_headers, table2_rows, "Table 2. Frozen case design and evidentiary limits."),
        ("Table_3_case_results.csv", table3_headers, table3_rows, "Table 3. Case-specific technical benchmark results."),
        ("Table_4_falsification.csv", table4_headers, table4_rows, "Table 4. Cross-case falsification and source-quality checks."),
        (
            "Table_S1_model_sensitivity.csv",
            ["Case", "Setting", "Estimate", "Unit"],
            sensitivity_rows,
            "Table S1. Pre-specified model sensitivity.",
        ),
    ]


def build_word_tables(specifications) -> None:
    document = configured_document()
    document.add_heading("Editable tables", 0)
    for _, headers, rows, caption in specifications:
        add_table(document, headers, rows, caption)
    document.save(TABLE_DIR / "Tables_editable.docx")


def build_pptx_tables(specifications) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for _, headers, rows, caption in specifications:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(
            Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6)
        )
        title.text_frame.text = caption
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        title.text_frame.paragraphs[0].runs[0].font.bold = True
        max_rows = min(len(rows), 13)
        table = slide.shapes.add_table(
            max_rows + 1,
            len(headers),
            Inches(0.25),
            Inches(0.9),
            Inches(12.8),
            Inches(6.2),
        ).table
        for column, header in enumerate(headers):
            cell = table.cell(0, column)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(217, 234, 247)
        for row_index, row in enumerate(rows[:max_rows], start=1):
            for column, value in enumerate(row):
                table.cell(row_index, column).text = str(value)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.LEFT
                    for run in paragraph.runs:
                        run.font.size = Pt(7 if len(headers) > 5 else 9)
                        run.font.name = "Arial"
    presentation.save(TABLE_DIR / "Tables_editable.pptx")


def main() -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    values = load_values()
    results = load_results()
    specifications = tables(values, results)
    for filename, headers, rows, _ in specifications:
        write_csv(filename, headers, rows)
    build_word_tables(specifications)
    build_pptx_tables(specifications)
    print(f"Wrote {len(specifications)} tables and editable DOCX/PPTX")


if __name__ == "__main__":
    main()
