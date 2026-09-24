import csv
import hashlib
import shutil
import subprocess
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from build_manuscript import TITLE
from publication_utils import OUTPUT, configured_document, display, load_values


SUBMISSION_DIR = OUTPUT / "submission"


def write_text_file(name: str, text: str) -> None:
    (SUBMISSION_DIR / name).write_text(text.rstrip() + "\n", encoding="utf-8")


def build_cover_letter(values) -> None:
    document = configured_document()
    document.add_paragraph("[Submission date]")
    document.add_paragraph("Editor-in-Chief")
    document.add_paragraph("Technological Forecasting and Social Change")
    document.add_paragraph()
    document.add_paragraph("Dear Editor,")
    document.add_paragraph(
        f"We submit the full-length research article “{TITLE}” for consideration in "
        "Technological Forecasting and Social Change."
    )
    document.add_paragraph(
        "The manuscript develops a pre-specified Legacy Standard Audit that separates "
        "historical, current frictionless, and transition-adjusted optima. Its contribution "
        "is methodological and governance-oriented: it makes claims of technology-driven "
        "mismatch falsifiable while preserving null, contrary, and insufficient-evidence "
        "outcomes."
    )
    document.add_paragraph(
        f"The empirical demonstration began from {display(values, 'DESIGN_CANDIDATES')} "
        f"candidates in {display(values, 'DESIGN_DOMAINS')} domains and froze "
        f"{display(values, 'DESIGN_FROZEN_CASES')} heterogeneous cases before confirmatory "
        "estimation. Results are deliberately not pooled because objectives and units are "
        "incomparable. The manuscript makes no general causal claim and withholds redesign "
        "recommendations where switching, compatibility, safety, or coordination costs are "
        "not observed."
    )
    document.add_paragraph(
        "The work fits the journal because specific technological developments—sensing, "
        "computation, communication, prognostics, interval metering, and adaptive control—"
        "are the central mechanisms that alter feasible decision sets. The paper connects "
        "these mechanisms to standards governance, transition costs, and policy implications."
    )
    document.add_paragraph(
        "The manuscript is original, is not under consideration elsewhere, and uses public "
        "secondary data and simulations. All authors must confirm authorship, conflicts, and "
        "the final declarations before submission."
    )
    document.add_paragraph("Sincerely,")
    document.add_paragraph("[Corresponding author name and affiliation]")
    path = SUBMISSION_DIR / "cover_letter.docx"
    document.save(path)
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(SUBMISSION_DIR),
            str(path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )


def build_graphical_abstract() -> None:
    fig = plt.figure(figsize=(13.28, 5.31), dpi=100)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(
        0.5,
        0.91,
        "Legacy Standard Audit",
        ha="center",
        va="center",
        fontsize=25,
        weight="bold",
        color="#17324D",
    )
    boxes = [
        (0.04, "Reconstruct", "objective • information\nconstraints • coordination"),
        (0.28, "Identify change", "sensing • computation\ncommunication • control"),
        (0.52, "Estimate", "static regret\nuncertainty • falsification"),
        (0.76, "Decide", "switching • network\nsafety • regulation"),
    ]
    for x, title, subtitle in boxes:
        ax.add_patch(
            plt.Rectangle(
                (x, 0.42),
                0.19,
                0.27,
                facecolor="#EAF4F8",
                edgecolor="#176B87",
                linewidth=2,
            )
        )
        ax.text(x + 0.095, 0.60, title, ha="center", va="center", weight="bold", fontsize=15)
        ax.text(x + 0.095, 0.49, subtitle, ha="center", va="center", fontsize=11)
        if x < 0.76:
            ax.annotate(
                "",
                xy=(x + 0.235, 0.555),
                xytext=(x + 0.195, 0.555),
                arrowprops={"arrowstyle": "->", "lw": 2, "color": "#6B7280"},
            )
    ax.text(
        0.5,
        0.20,
        "Technological change triggers re-audit—not automatic replacement.",
        ha="center",
        va="center",
        fontsize=18,
        weight="bold",
        color="#C8553D",
    )
    ax.text(
        0.5,
        0.08,
        "Retain material, uncertain, null, and insufficient-evidence outcomes.",
        ha="center",
        va="center",
        fontsize=13,
        color="#374151",
    )
    fig.savefig(
        SUBMISSION_DIR / "graphical_abstract.png",
        dpi=100,
        facecolor="white",
    )
    plt.close(fig)

    presentation = Presentation()
    presentation.slide_width = Inches(13.28)
    presentation.slide_height = Inches(5.31)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.15), Inches(12.28), Inches(0.6)
    )
    title_box.text_frame.text = "Legacy Standard Audit"
    title_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    title_box.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
    title_box.text_frame.paragraphs[0].runs[0].font.bold = True
    for index, (_, title, subtitle) in enumerate(boxes):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.45 + index * 3.2),
            Inches(1.55),
            Inches(2.65),
            Inches(1.45),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(234, 244, 248)
        shape.line.color.rgb = RGBColor(23, 107, 135)
        shape.text_frame.text = f"{title}\n{subtitle}"
        for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(17 if paragraph_index == 0 else 12)
                run.font.bold = paragraph_index == 0
        if index < 3:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(3.05 + index * 3.2),
                Inches(2.05),
                Inches(0.6),
                Inches(0.4),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(107, 114, 128)
            arrow.line.fill.background()
    footer = slide.shapes.add_textbox(
        Inches(0.8), Inches(3.75), Inches(11.7), Inches(0.8)
    )
    footer.text_frame.text = (
        "Technological change triggers re-audit—not automatic replacement.\n"
        "Retain material, uncertain, null, and insufficient-evidence outcomes."
    )
    footer.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    footer.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    footer.text_frame.paragraphs[0].runs[0].font.bold = True
    presentation.save(SUBMISSION_DIR / "graphical_abstract_editable.pptx")


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_manifest_and_zip() -> None:
    files = [
        OUTPUT / "manuscript" / "title_page.docx",
        OUTPUT / "manuscript" / "manuscript_anonymized.docx",
        OUTPUT / "manuscript" / "manuscript_anonymized.pdf",
        OUTPUT / "supplement" / "supplement.docx",
        OUTPUT / "supplement" / "supplement.pdf",
        OUTPUT / "figures" / "Figure_1_audit_framework.png",
        OUTPUT / "figures" / "Figure_2_primary_case_benchmarks.png",
        OUTPUT / "figures" / "Figure_3_case_classifications.png",
        OUTPUT / "figures" / "Figure_S1_model_sensitivity.png",
        OUTPUT / "figures" / "Figures_editable.pptx",
        OUTPUT / "tables" / "Tables_editable.docx",
        OUTPUT / "tables" / "Tables_editable.pptx",
        SUBMISSION_DIR / "cover_letter.docx",
        SUBMISSION_DIR / "cover_letter.pdf",
        SUBMISSION_DIR / "highlights.txt",
        SUBMISSION_DIR / "data_availability_statement.txt",
        SUBMISSION_DIR / "declarations.txt",
        SUBMISSION_DIR / "graphical_abstract.png",
        SUBMISSION_DIR / "graphical_abstract_editable.pptx",
    ]
    for path in files:
        if not path.exists():
            raise FileNotFoundError(path)
    manifest_path = SUBMISSION_DIR / "submission_manifest.csv"
    with manifest_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["file", "bytes", "sha256"])
        for path in files:
            writer.writerow([path.name, path.stat().st_size, file_hash(path)])
    archive_path = SUBMISSION_DIR / "TFSC_submission_package.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in files + [manifest_path]:
            archive.write(path, arcname=path.name)


def main() -> None:
    if SUBMISSION_DIR.exists():
        for path in SUBMISSION_DIR.iterdir():
            if path.is_file():
                path.unlink()
            elif path.is_dir():
                shutil.rmtree(path)
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    values = load_values()
    highlights = [
        "A frozen audit separates static mismatch from transition-adjusted value.",
        "Technology changes information and control constraints, not outcomes alone.",
        "Oracle benchmarks are retained without being treated as deployment effects.",
        "Null and insufficient-evidence cases prevent a forced redesign narrative.",
        "Heterogeneous case objectives are not pooled into a universal effect size.",
    ]
    if any(len(item) > 85 for item in highlights):
        raise ValueError("TFSC highlights must be no longer than 85 characters")
    write_text_file("highlights.txt", "\n".join(f"• {item}" for item in highlights))
    write_text_file(
        "data_availability_statement.txt",
        "Public data sources, retrieval metadata, local raw-snapshot paths, file sizes, "
        "SHA-256 hashes, and terms are listed in SOURCE_REGISTRY.csv. Analysis code, "
        "frozen case-selection records, and generated manuscript values accompany the "
        "submission repository. Copyrighted ISO sample material is not redistributed; "
        "its provenance and checksum are recorded.",
    )
    write_text_file(
        "declarations.txt",
        "Funding: No project-specific funding statement supplied.\n"
        "Competing interests: Authors must confirm before submission.\n"
        "Ethics: Public secondary data and simulations only; no participants recruited.\n"
        "Generative AI: Devin, built by Cognition AI, supported code generation, document "
        "assembly, and consistency checks. Human authors retain responsibility for all "
        "sources, analyses, interpretations, and final text.",
    )
    build_cover_letter(values)
    build_graphical_abstract()
    build_manifest_and_zip()
    print("Wrote TFSC submission assets and archive")


if __name__ == "__main__":
    main()
