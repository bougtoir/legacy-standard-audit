import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT / "outputs"


def load_values() -> dict[str, dict[str, str]]:
    with (PROJECT / "MANUSCRIPT_VALUES.csv").open(
        newline="", encoding="utf-8"
    ) as handle:
        return {row["value_id"]: row for row in csv.DictReader(handle)}


def load_results() -> dict[str, dict[str, object]]:
    payload = json.loads(
        (OUTPUT / "analysis" / "case_results.json").read_text(encoding="utf-8")
    )
    return {
        str(result["case_id"]): result for result in payload["case_results"]
    }


def numeric(values: dict[str, dict[str, str]], value_id: str) -> float:
    return float(values[value_id]["estimate"])


def display(values: dict[str, dict[str, str]], value_id: str) -> str:
    row = values[value_id]
    estimate = row["estimate"]
    spec = row["format_spec"]
    if spec == "s":
        return estimate.replace("_", " ")
    number = float(estimate)
    if spec.endswith("%"):
        return format(number, spec)
    if spec.endswith("d"):
        return format(int(round(number)), spec)
    return format(number, spec)


def interval(values: dict[str, dict[str, str]], value_id: str) -> str:
    row = values[value_id]
    spec = row["format_spec"]
    lower = float(row["lower"])
    upper = float(row["upper"])
    return f"{format(lower, spec)}–{format(upper, spec)}"


def configured_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 12)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return document


def add_page_number(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[object]],
    caption: str,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(caption)
    run.bold = True
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        shade_cell(table.rows[0].cells[index], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    document.add_paragraph()


def add_figure(
    document: Document,
    image_path: Path,
    caption: str,
    width_inches: float = 6.4,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = caption_paragraph.add_run(caption)
    run.bold = True


def add_landscape_section(document: Document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = 1
    section.page_width, section.page_height = section.page_height, section.page_width
    return section
